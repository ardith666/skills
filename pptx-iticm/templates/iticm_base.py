"""
pptx-iticm branding helper — V3 reference style.

Light content titles (dark 36pt + orange underline, NO navy header bar),
rounded cards (adjust 0.167, 1pt #E0E0E0), orange oval badges,
code in Consolas #C0C5D0, logo ONLY on title + closing slides.
Max 3 cards per slide — split longer lists across slides.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Brand Tokens ──────────────────────────────────────────
DARK_BG      = RGBColor(0x1a, 0x1a, 0x2e)
ACCENT       = RGBColor(0xe8, 0x6c, 0x00)
WHITE        = RGBColor(0xff, 0xff, 0xff)
LIGHT_GRAY   = RGBColor(0xf5, 0xf5, 0xf5)
MUTED        = RGBColor(0x6b, 0x70, 0x80)
DARK_TEXT    = RGBColor(0x1a, 0x1a, 0x2e)
CODE_BG      = RGBColor(0x2d, 0x2d, 0x3a)
CODE_LIGHT   = RGBColor(0xc0, 0xc5, 0xd0)
TASK_BG      = RGBColor(0x25, 0x25, 0x3a)
ANSWER_BAR   = RGBColor(0x28, 0xa7, 0x45)
CARD_BORDER  = RGBColor(0xe0, 0xe0, 0xe0)

FONT         = "Calibri"
CODE_FONT    = "Consolas"
DOSEN_NAME   = "Ardiyan NP S.Kom, M.Kom"
LABEL_URL    = "iticm.ac.id"
ROUND_ADJ    = 0.167

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(SKILL_DIR, "assets", "logo_iticm.png")


# ── Primitives ────────────────────────────────────────────
def _rrect(slide, left, top, width, height, color, line=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    try:
        shape.adjustments[0] = ROUND_ADJ
    except Exception:
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line:
        try:
            shape.line.color.rgb = CARD_BORDER
            shape.line.width = Pt(1)
        except Exception:
            pass
    else:
        shape.line.fill.background()
    return shape


def _oval(slide, left, top, diameter, color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(diameter), Inches(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color or ACCENT
    shape.line.fill.background()
    return shape


def _bar(slide, left, top, width=1.5, height=0.04):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    return shape


def _accent_edge(slide, left, top, height, width=0.06):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    return shape


def _set_bg(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height, text,
                 font_size=12, bold=False, color=WHITE,
                 alignment=PP_ALIGN.LEFT, font_name=FONT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def apply_branding(slide, show_name=False, dark_bg=True, logo=False):
    """Footer label (+ optional logo). Logo ONLY on title/closing (logo=True)."""
    label_color = WHITE if dark_bg else DARK_TEXT
    if logo and os.path.exists(LOGO_PATH):
        logo_size = Inches(1.0)
        slide.shapes.add_picture(
            LOGO_PATH,
            Inches(13.333 - 1.0 - 0.6),
            Inches(0.4),
            logo_size, logo_size)
    _bar(slide, left=0.7, top=7.0, width=2.0, height=0.03)
    _add_textbox(slide, 0.7, 7.08, 3.0, 0.3,
                 LABEL_URL, font_size=11, color=label_color)
    if show_name:
        _add_textbox(slide, 0.7, 6.55, 5.0, 0.4,
                     DOSEN_NAME, font_size=14, color=label_color)


def _light_base(prs, title, subtitle=None):
    """Light content slide: dark title + orange underline. Returns (slide, top)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, LIGHT_GRAY)
    _add_textbox(slide, 0.8, 0.5, 11.5, 0.7,
                 title, font_size=36, bold=True, color=DARK_TEXT)
    _bar(slide, 0.8, 1.15)
    if subtitle:
        _add_textbox(slide, 0.8, 1.35, 11.5, 0.45,
                     subtitle, font_size=18, color=MUTED)
        return slide, 2.0
    return slide, 1.6


def _code_textbox(slide, left, top, width, height, code,
                  font_size=13, color=None, max_lines=14):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code.split("\n")[:max_lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.name = CODE_FONT
        p.font.color.rgb = color or CODE_LIGHT
        p.space_after = Pt(1)
    return txBox


# ── Basic Slides ──────────────────────────────────────────
def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK_BG)
    _bar(slide, 0, 0, 13.333, 0.12)
    _add_textbox(slide, 1.5, 1.8, 10, 1.2,
                 title, font_size=54, bold=True, color=WHITE)
    if subtitle:
        _add_textbox(slide, 1.5, 3.2, 10, 0.8,
                     subtitle, font_size=22,
                     color=RGBColor(0xb0, 0xb5, 0xc0))
    _bar(slide, 1.5, 4.3, 2.5, 0.04)
    apply_branding(slide, show_name=True, dark_bg=True, logo=True)
    return slide


def add_content_slide(prs, title, items, note=""):
    """Generic bullets (fallback). For rich layouts use add_cards_slide etc."""
    slide, y = _light_base(prs, title)
    for item in items[:8]:
        if isinstance(item, tuple):
            heading, detail = item
            _add_textbox(slide, 1.0, y, 11.0, 0.4,
                         heading, font_size=20, bold=True, color=DARK_TEXT)
            y += 0.5
            _add_textbox(slide, 1.2, y, 10.8, 0.5,
                         detail, font_size=16, color=MUTED)
            y += 0.65
        else:
            _add_textbox(slide, 1.0, y, 11.0, 0.45,
                         f"•  {item}", font_size=18, color=DARK_TEXT)
            y += 0.6
    if note:
        _add_textbox(slide, 1.0, 6.5, 11.0, 0.3,
                     note, font_size=11, color=MUTED)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_two_column_slide(prs, title, left_items, right_items,
                         left_heading="", right_heading=""):
    slide, top = _light_base(prs, title)
    y = top
    if left_heading:
        _add_textbox(slide, 0.8, y, 5.5, 0.4,
                     left_heading, font_size=18, bold=True, color=DARK_TEXT)
        y += 0.5
    for item in left_items[:6]:
        _add_textbox(slide, 1.0, y, 5.3, 0.4,
                     f"•  {item}", font_size=16, color=DARK_TEXT)
        y += 0.55
    y = top
    if right_heading:
        _add_textbox(slide, 7.0, y, 5.5, 0.4,
                     right_heading, font_size=18, bold=True, color=DARK_TEXT)
        y += 0.5
    for item in right_items[:6]:
        _add_textbox(slide, 7.2, y, 5.3, 0.4,
                     f"•  {item}", font_size=16, color=DARK_TEXT)
        y += 0.55
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


# ── Rich Layouts (max 3 cards per slide — split longer lists) ──
def add_split_slide(prs, title, def_heading, def_detail,
                    ana_heading, ana_detail, badges):
    """Definition + analogy cards with orange edge + dark badge panel."""
    slide, _ = _light_base(prs, title)
    _rrect(slide, 0.8, 1.6, 7.5, 2.2, WHITE)
    _accent_edge(slide, 0.8, 1.6, 2.2)
    _add_textbox(slide, 1.2, 1.8, 6.8, 0.5,
                 def_heading, font_size=20, bold=True, color=DARK_TEXT)
    _add_textbox(slide, 1.2, 2.3, 6.8, 1.3,
                 def_detail, font_size=15, color=MUTED)
    _rrect(slide, 0.8, 4.1, 7.5, 1.8, WHITE)
    _accent_edge(slide, 0.8, 4.1, 1.8)
    _add_textbox(slide, 1.2, 4.3, 6.8, 0.5,
                 ana_heading, font_size=20, bold=True, color=DARK_TEXT)
    _add_textbox(slide, 1.2, 4.8, 6.8, 0.9,
                 ana_detail, font_size=15, color=MUTED)
    _rrect(slide, 8.8, 1.6, 3.8, 4.3, DARK_BG, line=False)
    y = 1.85
    for badge in badges[:5]:
        _add_textbox(slide, 9.2, y, 3.2, 0.6,
                     badge, font_size=15, color=WHITE)
        y += 0.78
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_cards_slide(prs, title, cards, subtitle=None):
    """Icon cards: [(emoji, heading, desc)] — 2 or 3 per call."""
    slide, top = _light_base(prs, title, subtitle)
    n = min(len(cards), 3)
    W, H, G = (3.8, 2.3, 0.3) if n == 3 else (4.2, 2.6, 0.4)
    X0 = 0.8 if n == 3 else 2.2
    for i, (emoji, head, desc) in enumerate(cards[:n]):
        x = X0 + i * (W + G)
        _rrect(slide, x, top, W, H, WHITE)
        _oval(slide, x + 0.3, top + 0.3, 0.6)
        _add_textbox(slide, x + 0.3, top + 0.35, 0.6, 0.5,
                     emoji, font_size=18, color=WHITE,
                     alignment=PP_ALIGN.CENTER)
        _add_textbox(slide, x + 1.1, top + 0.35, W - 1.3, 0.5,
                     head, font_size=18, bold=True, color=DARK_TEXT)
        _add_textbox(slide, x + 0.3, top + 1.1, W - 0.6, 1.1,
                     desc, font_size=14, color=MUTED)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_list_slide(prs, title, items, start=1, subtitle=None):
    """Numbered wide cards: [(heading, desc)] — max 3 per call."""
    slide, top = _light_base(prs, title, subtitle)
    y = top
    for k, (head, desc) in enumerate(items[:3]):
        _rrect(slide, 0.8, y, 11.7, 1.45, WHITE)
        _oval(slide, 1.1, y + 0.25, 0.55)
        _add_textbox(slide, 1.1, y + 0.3, 0.55, 0.45,
                     str(start + k), font_size=18, bold=True, color=WHITE,
                     alignment=PP_ALIGN.CENTER)
        _add_textbox(slide, 1.9, y + 0.2, 10.3, 0.5,
                     head, font_size=18, bold=True, color=DARK_TEXT)
        _add_textbox(slide, 1.9, y + 0.72, 10.3, 0.65,
                     desc, font_size=14, color=MUTED)
        y += 1.65
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_steps_slide(prs, title, steps, start=1, subtitle=None):
    """Big-number cards: [(heading, desc)] — max 3 per call."""
    slide, top = _light_base(prs, title, subtitle)
    W, H, G, X0 = 3.8, 2.6, 0.3, 0.8
    for i, (head, desc) in enumerate(steps[:3]):
        x = X0 + i * (W + G)
        _rrect(slide, x, top, W, H, WHITE)
        _oval(slide, x + W / 2 - 0.35, top + 0.2, 0.7)
        _add_textbox(slide, x + W / 2 - 0.35, top + 0.28, 0.7, 0.6,
                     str(start + i), font_size=22, bold=True, color=WHITE,
                     alignment=PP_ALIGN.CENTER)
        _add_textbox(slide, x + 0.3, top + 1.1, W - 0.6, 0.55,
                     head, font_size=17, bold=True, color=DARK_TEXT,
                     alignment=PP_ALIGN.CENTER)
        _add_textbox(slide, x + 0.3, top + 1.7, W - 0.6, 0.75,
                     desc, font_size=13, color=MUTED,
                     alignment=PP_ALIGN.CENTER)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_steps_code_slide(prs, title, subtitle, steps, code_title, code):
    """Left numbered steps (max 7) + right dark code panel."""
    slide, _ = _light_base(prs, title, subtitle)
    y = 2.1
    for i, step in enumerate(steps[:7]):
        _oval(slide, 1.0, y, 0.45)
        _add_textbox(slide, 1.0, y + 0.05, 0.45, 0.35,
                     str(i + 1), font_size=14, bold=True, color=WHITE,
                     alignment=PP_ALIGN.CENTER)
        _add_textbox(slide, 1.65, y + 0.03, 4.6, 0.45,
                     step, font_size=15, color=DARK_TEXT)
        y += 0.65
    _rrect(slide, 7.0, 1.6, 5.6, 4.8, DARK_BG, line=False)
    _add_textbox(slide, 7.3, 1.8, 5.0, 0.5,
                 code_title, font_size=20, bold=True, color=WHITE)
    _code_textbox(slide, 7.3, 2.4, 5.0, 3.8, code)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_explain_slide(prs, title, body, tip):
    """Detail card + dark tip bar."""
    slide, top = _light_base(prs, title)
    _rrect(slide, 0.8, top, 11.7, 3.1, WHITE)
    _accent_edge(slide, 0.8, top, 3.1)
    _add_textbox(slide, 1.2, top + 0.25, 10.9, 2.5,
                 body, font_size=16, color=DARK_TEXT)
    _rrect(slide, 0.8, top + 3.35, 11.7, 1.15, DARK_BG, line=False)
    _add_textbox(slide, 1.2, top + 3.55, 10.9, 0.8,
                 "💡 " + tip, font_size=15, color=WHITE)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


# ── Diagram Slide ─────────────────────────────────────────
def add_diagram_slide(prs, title, image_path, caption="",
                      subtitle="Diagram alur — baca dari atas ke bawah"):
    """Insert diagram PNG (flowchart via draw.io/mermaid → export PNG)."""
    import math
    slide, top = _light_base(prs, title, subtitle)
    if os.path.exists(image_path):
        try:
            from PIL import Image as _IMG
            with _IMG.open(image_path) as im:
                iw, ih = im.size
        except Exception:
            iw, ih = 1400, 1000
        maxw, maxh = 9.5, 4.6
        sc = min(maxw / (iw / 100.0), maxh / (ih / 100.0))
        w = (iw / 100.0) * sc
        h = (ih / 100.0) * sc
        x = (13.333 - w) / 2
        slide.shapes.add_picture(image_path, Inches(x), Inches(top + 0.1),
                                 Inches(w), Inches(h))
        if caption:
            _add_textbox(slide, 1.5, top + 0.1 + h + 0.1, 10.3, 0.35,
                         caption, font_size=12, color=MUTED,
                         alignment=PP_ALIGN.CENTER)
    elif caption:
        _add_textbox(slide, 1.5, top + 0.2, 10.3, 0.5,
                     "[diagram menyusul: " + caption + "]",
                     font_size=14, color=MUTED, alignment=PP_ALIGN.CENTER)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


# ── Code Slide ────────────────────────────────────────────
def add_code_slide(prs, title, code, lang="text"):
    """Full-width dark rounded code panel, Consolas CODE_LIGHT."""
    slide, top = _light_base(prs, title, lang.upper() if lang else None)
    _rrect(slide, 0.8, top, 11.7, 4.4, DARK_BG, line=False)
    _code_textbox(slide, 1.2, top + 0.2, 10.9, 4.0, code,
                  font_size=14, max_lines=13)
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


# ── References Slide ─────────────────────────────────────
def add_references_slide(prs, references):
    """Light-style bibliography."""
    slide, top = _light_base(prs, "Referensi")
    y = top
    for i, ref in enumerate(references):
        _add_textbox(slide, 1.0, y, 11.0, 0.5,
                     f"{i + 1}.  {ref}", font_size=14, color=DARK_TEXT)
        y += 0.6
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


# ── Summary Slide ─────────────────────────────────────────
def add_summary_slide(prs, title="Ringkasan", points=()):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK_BG)
    _bar(slide, 1.5, 0.35, 2.0, 0.12)
    _add_textbox(slide, 1.5, 0.55, 10, 0.7,
                 title, font_size=40, bold=True, color=WHITE)
    y = 1.6
    for point in points[:7]:
        _add_textbox(slide, 1.5, y, 10.3, 0.55,
                     "•  " + point, font_size=18, color=WHITE)
        y += 0.68
    apply_branding(slide, show_name=False, dark_bg=True)
    return slide


# ── Task Slides (OBE) ────────────────────────────────────
def add_example_task_slide(prs, title, examples):
    """Worked examples BEFORE the task. examples = [{"question","answer"}]."""
    slide, _ = _light_base(
        prs, title, "Contoh + pembahasan — pahami dulu sebelum tugas")
    y = 2.0
    for i, ex in enumerate(examples[:2]):
        _rrect(slide, 0.8, y, 11.7, 1.15, WHITE)
        _add_textbox(slide, 1.1, y + 0.12, 11.0, 0.45,
                     f"Soal {i + 1}: {ex.get('question', '')}",
                     font_size=16, bold=True, color=DARK_TEXT)
        _add_textbox(slide, 1.1, y + 0.58, 11.0, 0.5,
                     f"Jawaban: {ex.get('answer', '')}",
                     font_size=14, color=MUTED)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(y + 0.55),
            Inches(0.08), Inches(0.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ANSWER_BAR
        bar.line.fill.background()
        y += 1.35
    apply_branding(slide, show_name=False, dark_bg=False)
    return slide


def add_task_slide(prs, title, description, weight=""):
    """OBE task: dark bg + rounded card. Weight badge optional."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK_BG)
    _bar(slide, 1.5, 0.35, 2.5, 0.12)
    _add_textbox(slide, 1.5, 0.55, 9.0, 0.8,
                 title, font_size=38, bold=True, color=WHITE)
    if weight:
        _rrect(slide, 11.0, 0.6, 1.5, 0.5, ACCENT, line=False)
        _add_textbox(slide, 11.0, 0.62, 1.5, 0.45,
                     weight, font_size=14, bold=True, color=WHITE,
                     alignment=PP_ALIGN.CENTER)
    _rrect(slide, 1.2, 1.7, 10.9, 4.3, TASK_BG, line=False)
    _accent_edge(slide, 1.2, 1.7, 4.3)
    txBox = slide.shapes.add_textbox(
        Inches(1.7), Inches(2.0), Inches(10.0), Inches(3.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(description.split("\n")[:8]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(17)
        p.font.name = FONT
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)
    _add_textbox(slide, 1.5, 6.2, 10, 0.4,
                 "Kumpulkan sesuai tenggat yang ditentukan.",
                 font_size=13, color=MUTED)
    apply_branding(slide, show_name=False, dark_bg=True)
    return slide


def add_closing_slide(prs, text="Terima Kasih"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, DARK_BG)
    _bar(slide, 0, 0, 13.333, 0.12)
    _add_textbox(slide, 1.5, 2.8, 10, 1.0,
                 text, font_size=44, bold=True, color=ACCENT,
                 alignment=PP_ALIGN.CENTER)
    apply_branding(slide, show_name=True, dark_bg=True, logo=True)
    return slide
